"""Paint measured text lines as editable PowerPoint text."""

from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Pt

from pptx.oxml.ns import qn


_ALIGN = {
    'left': PP_ALIGN.LEFT,
    'right': PP_ALIGN.RIGHT,
    'center': PP_ALIGN.CENTER,
    'justify': PP_ALIGN.JUSTIFY,
    'start': None,
    'end': None,
}


def _is_rtl(style):
    direction = str((style or {}).get('direction') or '').lower()
    if direction in ('rtl', 'ltr'):
        return direction == 'rtl'
    return True


def _align_for(style, rtl):
    raw = str((style or {}).get('align') or '').lower()
    if raw == 'start':
        return PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    if raw == 'end':
        return PP_ALIGN.LEFT if rtl else PP_ALIGN.RIGHT
    return _ALIGN.get(raw) or (PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)


def _set_rtl(paragraph, rtl):
    try:
        paragraph._p.get_or_add_pPr().set('rtl', '1' if rtl else '0')
    except Exception:
        pass


def _set_spacing(run, px):
    if not px:
        return
    try:
        run._r.get_or_add_rPr().set('spc', str(_spacing_hundredths(px)))
    except Exception:
        pass


def _set_baseline(run, baseline):
    raw = str(baseline or '').lower()
    if raw in ('super', 'sub'):
        try:
            run._r.get_or_add_rPr().set(
                'baseline', '30000' if raw == 'super' else '-25000')
        except Exception:
            pass


def _paint_run(run, text, style, ctx):
    run.text = text
    try:
        run.font.name = ctx.font_name
    except Exception:
        pass
    try:
        run.font.size = _px_to_pt(style.get('size') or 14.0)
    except Exception:
        pass
    try:
        run.font.bold = int(style.get('weight') or 400) >= 600
        run.font.italic = bool(style.get('italic'))
    except Exception:
        pass
    rgb, alpha = _css_color(style.get('color'))
    if rgb is not None:
        try:
            run.font.color.rgb = rgb
        except Exception:
            pass
    combined = _combined_alpha(alpha, style.get('opacity'))
    _run_alpha(run, combined)
    try:
        if style.get('underline') is not None:
            run.font.underline = bool(style.get('underline'))
    except Exception:
        pass
    if style.get('strike'):
        try:
            run._r.get_or_add_rPr().set('strike', 'sngStrike')
        except Exception:
            pass
    _set_spacing(run, style.get('letter_spacing'))
    _set_baseline(run, style.get('baseline'))


def paint_text(slide, item, ctx):
    text = str(item.get('text') or '')
    if not text.strip():
        return 0
    style = item.get('style') or {}
    rtl = _is_rtl(style)
    x, y = float(item['x']), float(item['y'])
    w = max(1.0, float(item['w']))
    h = max(1.0, float(item['h']))
    # A line box measured by Chromium is already the exact glyph run: no wrap,
    # no autofit, no margins. This is what keeps the PPTX on the same pixels.
    box = slide.shapes.add_textbox(_px_to_emu(x), _px_to_emu(y),
                                   _px_to_emu(w + 2.0), _px_to_emu(h))
    frame = box.text_frame
    try:
        frame.word_wrap = False
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.margin_left = 0
        frame.margin_right = 0
        frame.margin_top = 0
        frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    except Exception:
        pass
    paragraph = frame.paragraphs[0]
    paragraph.alignment = _align_for(style, rtl)
    _set_rtl(paragraph, rtl)
    try:
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        paragraph.line_spacing = 1.0
    except Exception:
        pass
    run = paragraph.add_run()
    _paint_run(run, text, style, ctx)
    _no_line(box)
    ctx.texts.append(text)
    return 1


def paint_svg_text(slide, item, ctx):
    text = str(item.get('text') or '')
    if not text.strip():
        return 0
    x, y = float(item['x']), float(item['y'])
    w = max(1.0, float(item['w']))
    h = max(1.0, float(item['h']))
    style = {
        'size': item.get('size') or 10.0,
        'weight': item.get('weight') or 400,
        'italic': bool(item.get('italic')),
        'color': item.get('color'),
        'opacity': item.get('opacity'),
        'direction': item.get('direction') or 'ltr',
        'align': {'middle': 'center', 'end': 'right', 'start': 'left'}
                 .get(str(item.get('anchor') or 'start').lower(), 'left'),
    }
    return paint_text(slide, {'type': 'text', 'text': text, 'x': x, 'y': y,
                              'w': w, 'h': h, 'style': style}, ctx)
