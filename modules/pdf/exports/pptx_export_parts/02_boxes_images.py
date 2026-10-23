"""Paint measured boxes (cards, borders, shadows) and images."""

import base64
import math
from io import BytesIO

from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Emu


def _combined_alpha(css_alpha, opacity):
    """CSS fill transparency + element opacity -> one transparency percentage."""
    opacity = 1.0 if opacity is None else max(0.0, min(1.0, float(opacity)))
    base = 0.0 if css_alpha is None else max(0.0, min(100.0, float(css_alpha)))
    return 100.0 - ((100.0 - base) * opacity)


def _radius_px(token, width_px, height_px):
    text = str(token or '').strip()
    if not text:
        return 0.0
    first = text.split('/')[0].split()[0]
    if first.endswith('%'):
        try:
            percent = float(first[:-1])
        except ValueError:
            return 0.0
        return -1.0 if percent >= 50.0 else min(width_px, height_px) * percent / 100.0
    return _parse_px(first, 0.0) or 0.0


def _fill_shape(shape, item, width_px, height_px):
    opacity = item.get('opacity')
    gradients = item.get('gradients') or []
    gradient = gradients[0] if gradients else None
    if gradient:
        applied = _set_shape_gradient(shape, gradient)
        if applied:
            if opacity is not None and opacity < 0.999:
                _set_shape_alpha(shape, (1.0 - opacity) * 100.0)
            return
    rgb, alpha = _css_color(item.get('fill'))
    if rgb is not None:
        _set_shape_fill(shape, rgb, _combined_alpha(alpha, opacity))
    else:
        try:
            shape.fill.background()
        except Exception:
            pass


def _paint_borders(slide, item, x, y, w, h, shape):
    borders = [border for border in (item.get('borders') or [])
               if border.get('width', 0) > 0]
    if not borders:
        _no_line(shape)
        return
    uniform = len(borders) == 4 and len({(round(b['width'], 2), b['color'])
                                         for b in borders}) == 1
    if uniform:
        rgb, alpha = _css_color(borders[0]['color'])
        _set_shape_line(shape, rgb, borders[0]['width'],
                        dash=borders[0].get('style'),
                        alpha_pct=_combined_alpha(alpha, item.get('opacity')))
        return
    _no_line(shape)
    for border in borders:
        side = border['side']
        width = max(0.5, float(border['width']))
        rgb, alpha = _css_color(border['color'])
        if rgb is None:
            continue
        if side == 'top':
            bx, by, bw, bh = x, y, w, width
        elif side == 'bottom':
            bx, by, bw, bh = x, y + h - width, w, width
        elif side == 'left':
            bx, by, bw, bh = x, y, width, h
        else:
            bx, by, bw, bh = x + w - width, y, width, h
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, _px_to_emu(bx),
                                     _px_to_emu(by), _px_to_emu(max(bw, 0.1)),
                                     _px_to_emu(max(bh, 0.1)))
        _strip_theme_style(bar)
        _set_shape_fill(bar, rgb, _combined_alpha(alpha, item.get('opacity')))
        _no_line(bar)


def paint_box(slide, item, ctx):
    x, y = float(item['x']), float(item['y'])
    w, h = max(0.1, float(item['w'])), max(0.1, float(item['h']))
    radius = _radius_px(item.get('radius'), w, h)
    shape = _shape_for_rect(slide, x, y, w, h, radius)
    _fill_shape(shape, item, w, h)
    _paint_borders(slide, item, x, y, w, h, shape)
    _add_outer_shadow(shape, item.get('shadow'))
    return 1


def _natural_size(item, data):
    natural = item.get('natural') or {}
    if natural.get('w') and natural.get('h'):
        return float(natural['w']), float(natural['h'])
    return _image_size(data)


def _fit_mode(item):
    fit = str(item.get('fit') or 'fill').strip().lower()
    if fit.startswith('cover') or fit in ('100% 100%', '100%'):
        return 'cover' if fit.startswith('cover') else 'fill'
    if fit.startswith('contain') or fit == 'scale-down':
        return 'contain'
    if fit == 'none':
        return 'none'
    if fit == 'fill':
        return 'fill'
    if fit == 'auto':
        return 'auto'
    return 'explicit'


def _explicit_size(item, natural):
    fit = str(item.get('fit') or '').strip().lower()
    parts = [part for part in fit.split() if part]
    if not parts:
        return None
    width = _parse_px(parts[0], None)
    height = _parse_px(parts[1], None) if len(parts) > 1 else None
    if width is None:
        return None
    if width >= 0 and parts[0].endswith('%'):
        return None
    if height is None:
        height = width * natural[1] / natural[0] if natural[0] else width
    return width, height


def _background_position(item, box_w, box_h, draw_w, draw_h):
    raw = str(item.get('position') or '0% 0%').strip().lower()
    tokens = [token for token in raw.replace(',', ' ').split() if token]
    if len(tokens) == 1:
        token = tokens[0]
        if token in ('center', 'top', 'bottom'):
            return (box_w - draw_w) / 2.0, (box_h - draw_h) / 2.0 if token == 'center' else 0.0
        if token in ('left', 'right'):
            return (0.0 if token == 'left' else box_w - draw_w), (box_h - draw_h) / 2.0
        tokens = [token, token]
    def axis(token, box, draw):
        token = token.strip()
        if token == 'left' or token == 'top':
            return 0.0
        if token == 'right' or token == 'bottom':
            return box - draw
        if token == 'center':
            return (box - draw) / 2.0
        if token.endswith('%'):
            try:
                return (box - draw) * float(token[:-1]) / 100.0
            except ValueError:
                return 0.0
        return _parse_px(token, 0.0) or 0.0
    return axis(tokens[0], box_w, draw_w), axis(tokens[1], box_h, draw_h)


def _cover_crop(natural, box_w, box_h):
    image_ratio = natural[0] / natural[1] if natural[1] else 1.0
    box_ratio = box_w / box_h if box_h else 1.0
    left = right = top = bottom = 0.0
    if image_ratio > box_ratio:
        visible = box_ratio / image_ratio
        left = right = (1.0 - visible) / 2.0
    elif image_ratio < box_ratio:
        visible = image_ratio / box_ratio
        top = bottom = (1.0 - visible) / 2.0
    return left, top, right, bottom


def paint_image(slide, item, ctx):
    data = _image_bytes(item.get('src'), ctx.tenant_id)
    if not data:
        ctx.missing_images.append(item.get('src', '')[:300])
        return 0
    natural = _natural_size(item, data)
    if not natural[0] or not natural[1]:
        natural = _image_size(data)
    x, y = float(item['x']), float(item['y'])
    w, h = max(1.0, float(item['w'])), max(1.0, float(item['h']))
    mode = _fit_mode(item)
    opacity = item.get('opacity')
    radius = _radius_px(item.get('radius'), w, h)

    crop = None
    draw_rect = (x, y, w, h)
    if mode == 'cover':
        crop = _cover_crop(natural, w, h)
    elif mode == 'contain':
        scale = min(w / natural[0], h / natural[1])
        dw, dh = natural[0] * scale, natural[1] * scale
        draw_rect = (x + (w - dw) / 2.0, y + (h - dh) / 2.0, dw, dh)
    elif mode == 'none':
        dw, dh = natural
        draw_rect = (x + (w - dw) / 2.0, y + (h - dh) / 2.0, dw, dh)
    elif mode == 'auto' and item.get('background'):
        dw, dh = natural
        ox, oy = _background_position(item, w, h, dw, dh)
        if dw > w or dh > h:
            crop = _cover_crop(natural, w, h)
        else:
            draw_rect = (x + ox, y + oy, dw, dh)
    elif mode in ('explicit', 'auto') and item.get('background'):
        explicit = _explicit_size(item, natural)
        if explicit:
            dw, dh = explicit
            ox, oy = _background_position(item, w, h, dw, dh)
            draw_rect = (x + ox, y + oy, dw, dh)
    elif mode == 'explicit' and not item.get('background'):
        draw_rect = (x, y, w, h)

    try:
        picture = slide.shapes.add_picture(
            BytesIO(data), _px_to_emu(draw_rect[0]), _px_to_emu(draw_rect[1]),
            _px_to_emu(max(1.0, draw_rect[2])), _px_to_emu(max(1.0, draw_rect[3])))
    except Exception:
        ctx.missing_images.append(item.get('src', '')[:300])
        return 0
    if crop:
        try:
            picture.crop_left = crop[0]
            picture.crop_top = crop[1]
            picture.crop_right = crop[2]
            picture.crop_bottom = crop[3]
        except Exception:
            pass
    _set_picture_alpha(picture, None if opacity is None else (1.0 - opacity) * 100.0)
    _picture_shape(picture, draw_rect[2], draw_rect[3], radius)
    return 1


def paint_raster(slide, item, ctx):
    try:
        data = base64.b64decode(item.get('raster') or '')
    except Exception:
        return 0
    if not data:
        return 0
    slide.shapes.add_picture(BytesIO(data), _px_to_emu(0), _px_to_emu(0),
                             _px_to_emu(SLIDE_W_PX), _px_to_emu(SLIDE_H_PX))
    return 1
