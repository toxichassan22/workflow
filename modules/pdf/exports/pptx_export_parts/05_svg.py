"""Paint measured SVG primitives as native PowerPoint shapes."""

from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE


def _svg_dash(style):
    raw = str(style or '').lower()
    if 'dash' in raw:
        return 'dashed'
    if 'dot' in raw:
        return 'dotted'
    return None


def _apply_svg_paint(shape, item):
    opacity = item.get('opacity')
    gradient = item.get('gradient')
    applied = False
    if gradient:
        applied = _set_shape_gradient(shape, gradient)
    if not applied:
        rgb, alpha = _css_color(item.get('fill'))
        if rgb is not None:
            _set_shape_fill(shape, rgb, _combined_alpha(alpha, opacity))
        else:
            try:
                shape.fill.background()
            except Exception:
                pass
    elif opacity is not None and opacity < 0.999:
        _set_shape_alpha(shape, (1.0 - opacity) * 100.0)
    stroke = item.get('stroke')
    width = item.get('stroke_width') or 0
    rgb, alpha = _css_color(stroke)
    if rgb is not None and width > 0:
        _set_shape_line(shape, rgb, width, dash=_svg_dash(item.get('dash')),
                        alpha_pct=_combined_alpha(alpha, opacity))
    else:
        _no_line(shape)


def paint_svg_rect(slide, item, ctx):
    x, y = float(item['x']), float(item['y'])
    w, h = max(0.1, float(item['w'])), max(0.1, float(item['h']))
    radius = float(item.get('radius') or 0)
    shape = _shape_for_rect(slide, x, y, w, h, radius)
    _apply_svg_paint(shape, item)
    return 1


def paint_svg_ellipse(slide, item, ctx):
    x, y = float(item['x']), float(item['y'])
    w, h = max(0.1, float(item['w'])), max(0.1, float(item['h']))
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, _px_to_emu(x), _px_to_emu(y),
                                   _px_to_emu(w), _px_to_emu(h))
    _strip_theme_style(shape)
    _apply_svg_paint(shape, item)
    return 1


def paint_svg_line(slide, item, ctx):
    x1, y1 = float(item['x1']), float(item['y1'])
    x2, y2 = float(item['x2']), float(item['y2'])
    connector = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, _px_to_emu(x1), _px_to_emu(y1),
        _px_to_emu(x2), _px_to_emu(y2))
    _strip_theme_style(connector)
    opacity = item.get('opacity')
    rgb, alpha = _css_color(item.get('stroke'))
    width = item.get('stroke_width') or 1.0
    _set_shape_line(connector, rgb, width, dash=_svg_dash(item.get('dash')),
                    alpha_pct=_combined_alpha(alpha, opacity))
    return 1


def paint_svg_polyline(slide, item, ctx):
    points = [(float(x), float(y)) for x, y in (item.get('points') or [])]
    if len(points) < 2:
        return 0
    closed = bool(item.get('closed'))
    try:
        builder = slide.shapes.build_freeform(
            points[0][0], points[0][1], scale=EMU_PER_PX)
        builder.add_line_segments(points[1:], close=closed)
        shape = builder.convert_to_shape(origin_x=0, origin_y=0)
    except Exception:
        ctx.missing.append(('svg-polyline', ''))
        return 0
    _apply_svg_paint(shape, item)
    return 1
