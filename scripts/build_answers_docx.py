# -*- coding: utf-8 -*-
"""Convert clean_answers_reports.md to an Arabic RTL .docx.

Keeps: headings, bullets, numbered items, blockquotes, tables, bold/italic/code,
and <span style="color:red"> highlights -> red runs.
"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\workflow\clean_answers_reports.md"
OUT = r"D:\workflow\أجوبة_التقارير.docx"

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GRAY = RGBColor(0x59, 0x59, 0x59)
RED = RGBColor(0xB9, 0x1C, 0x1C)


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    bidi.set(qn("w:val"), "1")
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def style_run(run, size=12, bold=False, italic=False, color=None, code=False):
    run.font.name = "Consolas" if code else "Arial"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:cs"), "Consolas" if code else "Arial")
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(size * 2))
    rPr.append(sz_cs)
    run.font.rtl = True


TOKEN_RE = re.compile(
    r'(<span style="color:red">.*?</span>|\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)',
    re.DOTALL,
)


def add_runs(paragraph, text, size=12, force_red=False):
    for tok in TOKEN_RE.split(text):
        if not tok:
            continue
        red = force_red
        bold = italic = code = False
        if tok.startswith('<span style="color:red">'):
            tok = tok[len('<span style="color:red">'):-len('</span>')]
            red = True
        elif tok.startswith('**') and tok.endswith('**'):
            tok = tok[2:-2]
            bold = True
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            tok = tok[1:-1]
            italic = True
        elif tok.startswith('`') and tok.endswith('`'):
            tok = tok[1:-1]
            code = True
        # nested markdown inside a red span
        inner = add_runs(paragraph, tok, size=size, force_red=red) if red and TOKEN_RE.search(tok) else None
        if inner:
            continue
        run = paragraph.add_run(tok)
        style_run(run, size=size, bold=bold, italic=italic, code=code,
                  color=RED if red else None)
    return True


def add_par(doc, text, size=12, bold=False, color=None, indent=0,
            space_after=6, style=None):
    p = doc.add_paragraph(style=style)
    set_rtl(p)
    p.paragraph_format.space_after = Pt(space_after)
    if indent:
        p.paragraph_format.right_indent = Cm(indent)
    add_runs(p, text, size=size)
    if bold and not p.runs:
        p.add_run(text)
    for r in p.runs:
        if bold:
            r.font.bold = True
        if color is not None:
            r.font.color.rgb = color
    return p


def add_rule(doc):
    p = doc.add_paragraph()
    set_rtl(p)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "C8C8C8")
    pPr.append(pBdr)
    pPr.append(bottom)
    p.paragraph_format.space_after = Pt(10)


def set_table_rtl(table):
    table.alignment = WD_TABLE_ALIGNMENT.RIGHT
    tblPr = table._tbl.tblPr
    bidi = OxmlElement("w:bidiVisual")
    tblPr.append(bidi)


def add_table(doc, rows):
    header = rows[0]
    body = rows[2:] if len(rows) > 2 and set(rows[1][0].strip()) <= set('-: ') else rows[1:]
    table = doc.add_table(rows=0, cols=len(header))
    table.style = 'Table Grid'
    set_table_rtl(table)
    for r_i, cells in enumerate([header] + body):
        row = table.add_row()
        for c_i, cell_text in enumerate(cells):
            cell = row.cells[c_i]
            p = cell.paragraphs[0]
            set_rtl(p)
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, cell_text.strip(), size=10)
            for r in p.runs:
                if r_i == 0:
                    r.font.bold = True
    doc.add_paragraph()


def main():
    doc = Document()
    for section in doc.sections:
        section.right_margin = Cm(2.2)
        section.left_margin = Cm(2.2)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(12)
    style.element.rPr.rFonts.set(qn("w:cs"), "Arial")

    lines = open(SRC, encoding='utf-8').read().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue
        # table block
        if stripped.startswith('|') and stripped.endswith('|'):
            rows = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                cells = [c for c in lines[i].strip().strip('|').split('|')]
                rows.append(cells)
                i += 1
            add_table(doc, rows)
            continue
        if stripped == '---':
            add_rule(doc)
        elif stripped.startswith('### '):
            add_par(doc, stripped[4:], size=14, bold=True, color=NAVY, space_after=6)
        elif stripped.startswith('## '):
            add_par(doc, stripped[3:], size=16, bold=True, color=NAVY, space_after=8)
        elif stripped.startswith('# '):
            add_par(doc, stripped[2:], size=20, bold=True, color=NAVY, space_after=4)
        elif stripped.startswith('>'):
            add_par(doc, stripped.lstrip('> ').strip(), size=11, indent=0.8, color=GRAY,
                    space_after=6)
        elif stripped.startswith('- '):
            add_par(doc, stripped[2:], indent=0.6)
        elif re.match(r'^\d+\.\s', stripped):
            add_par(doc, stripped, indent=0.6)
        else:
            add_par(doc, stripped)
        i += 1

    doc.save(OUT)
    print('saved:', OUT)


if __name__ == '__main__':
    main()
