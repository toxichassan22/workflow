import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from exports.pptx_export import _verify_pptx_slide_count, generate_pptx


def _slide(title, text):
    return {
        'title': title,
        'type': 'content',
        'html': (
            '<div class="slide" style="width:1280px;height:720px;background:#ffffff;">'
            f'<h2>{title}</h2><p>{text}</p></div>'
        ),
    }


class PptxExportVerificationTests(unittest.TestCase):
    def test_generate_pptx_writes_every_slide(self):
        from pptx import Presentation

        slides = [_slide(f'عنوان {index}', f'نص الشريحة {index}') for index in range(3)]
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = generate_pptx(slides, 'عرض الاختبار', {}, tmp_dir, 'tenant-test')
            self.assertTrue(os.path.isfile(path))
            self.assertEqual(len(Presentation(path).slides), 3)

    def test_generate_pptx_rejects_empty_deck(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            with self.assertRaises(ValueError):
                generate_pptx([], 'عرض فارغ', {}, tmp_dir, 'tenant-test')

    def test_verify_pptx_slide_count_refuses_short_file(self):
        from pptx import Presentation

        slides = [_slide(f'عنوان {index}', f'نص {index}') for index in range(2)]
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = generate_pptx(slides, 'عرض التحقق', {}, tmp_dir, 'tenant-test')
            self.assertEqual(_verify_pptx_slide_count(path, 2), 2)
            with self.assertRaises(RuntimeError) as raised:
                _verify_pptx_slide_count(path, 3)
            self.assertIn('2', str(raised.exception))
            self.assertIn('3', str(raised.exception))
            self.assertEqual(len(Presentation(path).slides), 2)


if __name__ == '__main__':
    unittest.main()
