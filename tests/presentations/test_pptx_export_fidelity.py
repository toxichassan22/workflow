import base64
import os
import struct
import sys
import tempfile
import unittest
import zipfile

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..'))

from exports.pptx_export import generate_pptx, _verify_no_lost_text

_LOGO_PNG = (
    'iVBORw0KGgoAAAANSUhEUgAAAFAAAAAoCAYAAABpYH0BAAAA6UlEQVR4nO2aQQ7CIBQFS7f+O+gd9NLu'
    'PIiH8BDu27SxDZCYgi8RAjM7mq4m81jhpmkafOxyCz9AwPv1dP559A/IOyZ2tAtEXjq+q1Ug8vLZ'
    'nLnT+cqdJxDcgZAPAkUQKIJAEQTWJtDuj6EnKLAmgfapr6cKKbAWgRZV10uFFCiCwBoE2pe59jBj'
    'Ciwt0A4qa71CChRBYEmBljjPlmdMgaUE5lZljVZIgSIILCHw1zlagzOmwH8LVCuyxiqkQBGedohQ'
    'oAgCRRAogkCRMX7zC+ks7tYCkZjP5myfMBLT8V0FdyAShyx5y88zBE8++oCRw3IAAAAASUVORK5C'
    'YII='
)


def _slide_html():
    logo = 'data:image/png;base64,' + _LOGO_PNG
    return f'''
<div class="slide" dir="rtl" style="width:1280px;height:720px;position:relative;overflow:hidden;box-sizing:border-box;background:#ffffff;color:#0b1f33;font-family:Arial;">
  <div style="position:absolute;left:100px;top:100px;width:300px;height:120px;background:#f8fafc;border:1px solid #e2e8f0;border-radius:8px;box-shadow:0 1px 3px rgba(0,0,0,0.08);"></div>
  <img src="{logo}" style="position:absolute;left:1170px;top:20px;width:80px;height:40px;object-fit:fill;" />
  <h2 style="position:absolute;left:100px;top:40px;width:600px;margin:0;font-size:24px;font-weight:700;color:#0b1f33;">تحليل الحساسية العام</h2>
  <table style="position:absolute;left:100px;top:240px;width:600px;border-collapse:collapse;font-size:12px;">
    <thead><tr>
      <th style="background:#0b1f33;color:#ffffff;padding:6px;text-align:center;">البند</th>
      <th style="background:#0b1f33;color:#ffffff;padding:6px;text-align:center;">القيمة</th>
    </tr></thead>
    <tbody>
      <tr><td style="padding:5px;text-align:right;">إجمالي الإيرادات</td>
          <td style="padding:5px;text-align:center;"><span style="background:#eff6ff;border:1px solid #bfdbfe;border-radius:10px;padding:2px 8px;color:#1d4ed8;font-size:9px;">الأعلى أفضل</span></td></tr>
      <tr><td style="padding:5px;text-align:right;">إجمالي التكلفة</td><td style="padding:5px;text-align:center;">529,227,650.96</td></tr>
    </tbody>
  </table>
  <svg width="420" height="200" viewBox="0 0 420 200" style="position:absolute;left:760px;top:240px;">
    <rect x="0" y="0" width="420" height="200" fill="#ffffff" rx="6"/>
    <line x1="40" y1="160" x2="400" y2="160" stroke="#64748b" stroke-width="1.5"/>
    <rect x="60" y="100" width="40" height="60" fill="#0ea5e9" rx="2"/>
    <rect x="130" y="60" width="40" height="100" fill="#0b6e63" rx="2"/>
    <polyline points="90 120 160 80 230 60 300 40" fill="none" stroke="#0284c7" stroke-width="3"/>
    <circle cx="90" cy="120" r="4" fill="#ffffff" stroke="#0284c7" stroke-width="2"/>
    <text x="70" y="145" font-size="10" fill="#64748b" text-anchor="middle">السنة 1</text>
    <text x="140" y="45" font-size="10" fill="#0b1f33" text-anchor="middle">470.99</text>
  </svg>
  <p style="position:absolute;left:100px;top:480px;width:600px;margin:0;font-size:14px;">النص العربي يجب أن يظهر كاملاً في الملف.</p>
</div>'''


def _walk(shapes):
    for shape in shapes:
        yield shape
        if getattr(shape, 'shape_type', None) == 6:  # group
            yield from _walk(shape.shapes)


def _slide_text(slide):
    parts = []
    for shape in _walk(slide.shapes):
        if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
            parts.append(shape.text_frame.text)
        if getattr(shape, 'has_table', False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
    return '\n'.join(parts)


class PptxExportFidelityTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.path = generate_pptx(
            [{'title': 'اختبار', 'type': 'content', 'html': _slide_html()}],
            'عرض الدقة', {'font_family': 'The Sans Arabic'}, cls.tmp.name, 'tenant-test')
        from pptx import Presentation
        cls.prs = Presentation(cls.path)
        cls.slide = cls.prs.slides[0]
        cls.text = _slide_text(cls.slide)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_every_text_is_present(self):
        for phrase in ('تحليل الحساسية العام', 'إجمالي الإيرادات', 'إجمالي التكلفة',
                       'الأعلى أفضل', 'السنة 1', '470.99',
                       'النص العربي يجب أن يظهر كاملاً في الملف.'):
            self.assertIn(phrase, self.text)

    def test_table_is_native_and_carries_cells(self):
        tables = [shape for shape in _walk(self.slide.shapes)
                  if getattr(shape, 'has_table', False) and shape.has_table]
        self.assertEqual(len(tables), 1)
        table = tables[0].table
        self.assertEqual(len(table.rows), 3)
        self.assertEqual(len(table.columns), 2)
        # An RTL HTML table starts at the right; PowerPoint's first column is
        # the leftmost, so the visual order mirrors the source order.
        self.assertIn('البند', table.cell(0, 1).text)
        self.assertIn('القيمة', table.cell(0, 0).text)
        self.assertIn('529,227,650.96', table.cell(2, 0).text)

    def test_chart_primitives_are_native_shapes(self):
        # The old exporter reduced an SVG chart to twelve loose labels. The
        # chart must now keep its bars, line, marker and labels.
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        types = [shape.shape_type for shape in _walk(self.slide.shapes)]
        self.assertIn(MSO_SHAPE_TYPE.FREEFORM, types)
        self.assertIn(MSO_SHAPE_TYPE.AUTO_SHAPE, types)
        self.assertIn('470.99', self.text)

    def test_header_logo_picture_keeps_its_measured_rect(self):
        pictures = [shape for shape in _walk(self.slide.shapes)
                    if shape.shape_type == 13]
        self.assertTrue(pictures)
        logo = min(pictures, key=lambda shape: shape.top)
        left_px = logo.left / 914400 * 96
        top_px = logo.top / 914400 * 96
        self.assertAlmostEqual(left_px, 1170.0, delta=2.0)
        self.assertAlmostEqual(top_px, 20.0, delta=2.0)

    def test_font_is_embedded_in_the_package(self):
        with zipfile.ZipFile(self.path) as archive:
            names = archive.namelist()
            fonts = [name for name in names if name.startswith('ppt/fonts/')]
            self.assertTrue(fonts, 'expected embedded font parts')
            payload = archive.read(fonts[0])
            # PowerPoint stores fonts as uncompressed EOT, not raw TTF. The
            # version is the one PowerPoint's own embedded-font parts carry.
            self.assertEqual(payload[8:12], b'\x02\x00\x02\x00')
            self.assertEqual(struct.unpack('<H', payload[34:36])[0], 0x504C)
            presentation = archive.read('ppt/presentation.xml').decode('utf-8')
            self.assertIn('embeddedFontLst', presentation)
            self.assertIn('fntdata', archive.read('[Content_Types].xml').decode('utf-8')
                          + ' '.join(fonts))

    def test_lost_text_guard_refuses_a_missing_line(self):
        display = {'items': [{'type': 'text', 'text': 'سطر موجود'}]}
        with self.assertRaises(RuntimeError):
            _verify_no_lost_text(display, ['شيء آخر'], 7)


class PptxRasterSafetyNetTests(unittest.TestCase):
    def test_unsupported_css_falls_back_to_one_picture(self):
        html = (
            '<div class="slide" dir="rtl" style="width:1280px;height:720px;'
            'position:relative;overflow:hidden;background:#ffffff;">'
            '<div style="filter:blur(2px);width:200px;height:80px;'
            'background:#0b1f33;color:#fff;font-size:20px;">نص</div></div>')
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = generate_pptx(
                [{'title': 'راستر', 'type': 'content', 'html': html}],
                'عرض راستر', {}, tmp_dir, 'tenant-test')
            from pptx import Presentation
            slide = Presentation(path).slides[0]
            self.assertEqual(len(slide.shapes), 1)
            self.assertEqual(slide.shapes[0].shape_type, 13)


if __name__ == '__main__':
    unittest.main()
