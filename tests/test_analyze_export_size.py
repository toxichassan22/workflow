import contextlib
import io
import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), '..', 'scripts'))

from analyze_export_size import main


def _png_bytes(color):
    from PIL import Image
    buf = io.BytesIO()
    Image.new('RGB', (400, 300), color).save(buf, format='PNG')
    return buf.getvalue()


class AnalyzeExportSizeTests(unittest.TestCase):
    def _run(self, *paths):
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            code = main(list(paths))
        self.assertEqual(code, 0)
        return out.getvalue()

    def test_pptx_reports_shared_media_once(self):
        from pptx import Presentation
        from pptx.util import Inches

        red = _png_bytes((200, 30, 30))
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = os.path.join(tmp_dir, 'deck.pptx')
            prs = Presentation()
            prs.slide_width = Inches(13.333)
            prs.slide_height = Inches(7.5)
            blank = prs.slide_layouts[6]
            for _ in range(2):
                slide = prs.slides.add_slide(blank)
                slide.shapes.add_picture(io.BytesIO(red), Inches(1), Inches(1))
            prs.save(path)
            report = self._run(path)
        self.assertIn('2 slides', report)
        self.assertIn('1 parts', report)
        self.assertIn('no duplicate media parts', report)

    def test_pdf_flags_content_duplicates_across_pages(self):
        import fitz

        red = _png_bytes((200, 30, 30))
        with tempfile.TemporaryDirectory() as tmp_dir:
            path = os.path.join(tmp_dir, 'deck.pdf')
            doc = fitz.open()
            for _ in range(2):
                page = doc.new_page(width=600, height=400)
                page.insert_image(page.rect, stream=red)
            doc.save(path)
            doc.close()
            report = self._run(path)
        self.assertIn('2 pages', report)
        self.assertIn('1 refs', report)
        self.assertIn('x2', report)


if __name__ == '__main__':
    unittest.main()
